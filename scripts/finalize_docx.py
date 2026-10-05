#!/usr/bin/env python3
"""Inject real page-numbered Contents / List of Figures / List of Tables.

LibreOffice headless conversion does not populate TOC/LOF/LOT fields, so the
PDF showed only 'Right-click to Update Field' placeholders. This script
derives pagination from the converted PDF itself (2-pass: any page shift
caused by the injected lists triggers a re-derive + re-inject until stable,
max 3 rounds) and replaces the three field paragraphs with static,
page-numbered lists. Idempotent within a fresh make_paper_pro.py build;
re-runs replace the previously injected blocks via BEGIN/END markers.

Pipeline: make_paper_pro.py -> add_static_index.py -> finalize_docx.py
(The final PDF is produced by the last conversion inside this script.)
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

REPO = Path(__file__).resolve().parent.parent
DOCX = REPO / "docx" / "green-code-interaction-benchmark.docx"
PDF = REPO / "docx" / "green-code-interaction-benchmark.pdf"
WIDTH_TWIPS = "9350"
MARK = "STATIC-LIST"
MAX_ROUNDS = 3


def convert():
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf",
         str(DOCX), "--outdir", str(DOCX.parent)],
        capture_output=True, text=True, timeout=180)
    return PDF


def pdf_pages(pdf):
    n = int(subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                           text=True, timeout=60).stdout.split("Pages:")[1]
            .split()[0])
    pages = []
    for i in range(1, n + 1):
        out = subprocess.run(
            ["pdftotext", "-layout", "-f", str(i), "-l", str(i),
             str(pdf), "-"], capture_output=True, text=True,
            timeout=60).stdout
        pages.append(re.sub(r"\s+", " ", out))
    return pages


def nospace(s):
    return re.sub(r"\s+", "", s)


def find_page(pages, key, start=1):
    k = nospace(key)
    for i in range(start - 1, len(pages)):
        if k in nospace(pages[i]):
            return i + 1
    return None


def collect(doc):
    """Return (toc_headings, fig_captions, tab_captions) in document order,
    skipping previously injected static blocks."""
    toc, figs, tabs = [], [], []
    in_static = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if MARK in t:
            if "BEGIN" in t:
                in_static = True
            elif "END" in t:
                in_static = False
            continue
        if in_static:
            continue
        st = p.style.name
        if st.startswith("Heading") and t and t not in (
                "Table of Contents", "Contents", "List of Figures",
                "List of Tables"):
            if int(st.split()[-1]) <= 2:
                toc.append(t)
        if not t:
            continue
        try:
            is_cap = (p.style.name == "Caption")
        except Exception:
            is_cap = False
        if t.startswith("Figure ") and ". " in t[:12]:
            figs.append(t)
        elif is_cap and t.startswith("Table ") and ". " in t[:11]:
            tabs.append(t)
    return toc, figs, tabs


def derive(pages, doc):
    toc, figs, tabs = collect(doc)
    # Body starts at section 1; front-matter static lists (which repeat every
    # heading/caption verbatim) sit before it, so body items must be searched
    # only at/after the body start to avoid self-matches.
    sec1 = find_page(pages, "first-order concern for datacentres") or 3
    toc_pages, fig_pages, tab_pages = [], [], []
    body_started = False
    for t in toc:
        if t.startswith("1 Introduction"):
            body_started = True
        key = t[:42]
        pg = find_page(pages, key, start=(sec1 if body_started else 1))
        toc_pages.append((t, pg or "?"))
    for t in figs:
        m = re.match(r"Figure\s*(\d+)\.\s*(.*)", t)
        num, title = (m.group(1), m.group(2)) if m else ("?", t)
        key = title[:30]
        fig_pages.append((f"Figure {num}. {title}", t,
                          find_page(pages, key[:22], start=sec1) or "?"))
    for t in tabs:
        m = re.match(r"Table\s*(\d+)\.\s*(.*)", t)
        num, title = (m.group(1), m.group(2)) if m else ("?", t)
        key = title[:30]
        tab_pages.append((f"Table {num}. {title}", t,
                          find_page(pages, key[:22], start=sec1) or "?"))
    return toc_pages, fig_pages, tab_pages


def clear_between(doc, begin_mark, end_mark):
    """Remove previously injected block (idempotency across runs)."""
    els = list(doc.element.body)
    kill, inside = [], False
    for el in els:
        if not el.tag.endswith("p"):
            continue
        t = "".join(n.text or "" for n in el.iter()
                    if n.tag.endswith("t"))
        if begin_mark in t:
            inside, kill = True, [el]
            continue
        if end_mark in t and inside:
            kill.append(el)
            inside = False
            for k in kill:
                doc.element.body.remove(k)
            kill = []
            continue
        if inside:
            kill.append(el)


def entry_para(doc, ref_el, text, page, size=9.5):
    p_el = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), WIDTH_TWIPS)
    tabs.append(tab)
    pPr.append(tabs)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    sz2 = OxmlElement("w:szCs")
    sz2.set(qn("w:val"), str(int(size * 2)))
    rPr = OxmlElement("w:rPr")
    rPr.append(sz)
    pPr.append(rPr)
    p_el.append(pPr)
    for chunk, bold in ((text + "\t", False), (str(page), False)):
        r_el, t_el = OxmlElement("w:r"), OxmlElement("w:t")
        t_el.set(qn("xml:space"), "preserve")
        t_el.text = chunk
        rp = OxmlElement("w:rPr")
        s2 = OxmlElement("w:sz")
        s2.set(qn("w:val"), str(int(size * 2)))
        rp.append(s2)
        r_el.append(rp)
        r_el.append(t_el)
        p_el.append(r_el)
    ref_el.addnext(p_el)
    return p_el


def marker_para(text):
    p_el = OxmlElement("w:p")
    r_el, t_el = OxmlElement("w:r"), OxmlElement("w:t")
    t_el.text = text
    rph = OxmlElement("w:rPr")
    rph.append(OxmlElement("w:vanish"))  # hidden marker, keeps body clean
    r_el.append(rph)
    r_el.append(t_el)
    p_el.append(r_el)
    return p_el


def hide_runs(p):
    """Remove visible placeholder text from a field paragraph while keeping
    the field codes (fldChar/instrText) intact for Word users. (The original
    generator placed the placeholder <w:t> inside <w:fldChar>, which
    LibreOffice renders regardless of the hidden flag.)"""
    WT = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    for t in p._element.iter():
        if t.tag == WT:
            t.text = ""


def inject(doc, toc_pages, fig_pages, tab_pages):
    # locate the three field paragraphs
    targets = {}
    for p in doc.paragraphs:
        instr = "".join(
            (n.text or "") for n in p._element.iter()
            if n.tag.endswith("instrText"))
        if "TOC " not in instr and "INDEX " not in instr:
            continue
        if "INDEX" in instr:
            # strip its placeholder text too (static index follows it)
            for t in p._element.iter():
                if t.tag.endswith("t"):
                    t.text = ""
            continue
        if '"Figure"' in instr:
            targets["lof"] = p
        elif '"Table"' in instr:
            targets["lot"] = p
        else:
            targets["toc"] = p
    for key in ("toc", "lof", "lot"):
        clear_between(doc, f"{MARK}-{key}-BEGIN", f"{MARK}-{key}-END")
    blocks = {
        "toc": [(t, pg) for t, pg in toc_pages],
        "lof": [(full, pg) for full, _, pg in fig_pages],
        "lot": [(full, pg) for full, _, pg in tab_pages],
    }
    for key, entries in blocks.items():
        p = targets.get(key)
        if p is None:
            print(f"[finalize] WARNING: no field paragraph for {key}")
            continue
        hide_runs(p)  # hide placeholder; field codes survive for Word
        anchor = p._element
        anchor.addnext(marker_para(f"{MARK}-{key}-BEGIN"))
        cur = anchor.getnext()
        for text, pg in entries:
            cur = entry_para(doc, cur, text, pg)
        cur.addnext(marker_para(
            f"{MARK}-{key}-END static snapshot from this build's "
            f"pagination (Word: delete block + F9 to refresh)"))
    return True


def main():
    prev = None
    convert()  # fresh PDF from the field-placeholder build
    for rnd in range(1, MAX_ROUNDS + 1):
        doc = Document(DOCX)
        pages = pdf_pages(PDF)
        toc_pages, fig_pages, tab_pages = derive(pages, doc)
        sig = (str([(t, pg) for t, pg in toc_pages]),
               str([(f, pg) for f, _, pg in fig_pages]),
               str([(t, pg) for t, _, pg in tab_pages]))
        print(f"[finalize] round {rnd}: {len(toc_pages)} toc, "
              f"{len(fig_pages)} figs, {len(tab_pages)} tabs; "
              f"stable={sig == prev}")
        if sig == prev:
            break
        prev = sig
        inject(doc, toc_pages, fig_pages, tab_pages)
        doc.save(DOCX)
        convert()
    print(f"[finalize] wrote {PDF.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
